# -*- coding: utf-8 -*-
"""把「资料」文件夹里的文档转成 知识库.js —— 网页刷新后自动出现这些知识库。

用法（在站点目录里）：
    python make-kb.py                 扫 站点目录\\资料\\
    python make-kb.py D:\\书 D:\\讲义     每个目录 = 一个知识库（可给多个）

支持：.txt .md .csv .json .html .docx .pdf
      .docx 需要 python-docx（没有也能读：会退回直接解压的方式）
      .pdf  需要 pypdf（pip install pypdf），没装就跳过并提示

页码：
    · PDF 有书签 → 按书签切小节，每节记起始页（page / pages）；
    · PDF 没书签 → 看文件名里有没有页码（如「1.结核_110-117.pdf」），有就记下起始页。
      病理那 27 个精编版就是这种：本身没书签，但文件名写了它在原书第几页。

默认「合并 + 增量」：
    · 合并：读现有 知识库.js，同名库（按 id）被替换、其它库原样保留 —— 不再整份覆盖丢库。
    · 增量：每个文件算 md5 存进 知识库.清单.json，没变的文件直接复用旧正文、不重新解析，
      大 PDF 第二次跑会快很多。
    · 书签切分：PDF 带书签时按顶层书签切成小节（每节 = 一篇），避免「一本书 = 一篇」
      在检索时被「同一文件最多 2 段」卡住。

想回到旧行为（整份覆盖、不合并、全部重解析）加 --overwrite。
可选参数：
    --overwrite      不合并现有库、不读缓存，只输出本次扫描的库（旧的整份覆盖行为）
    --skip-pdf       跳过 PDF（最快）
    --no-bookmarks   关闭 PDF 书签切分（一本书仍算一篇）
    --pdf-pages 30   每个 PDF 只读前 30 页
    --max-chars 60000 每篇最多留多少字
    --dry            只统计、不写文件
    -h, --help       看这段说明
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding='utf-8')   # Windows 控制台里中文才不乱码
except Exception:
    pass
import json
import io
import re
import zipfile
import hashlib

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, '知识库.js')
MANIFEST = os.path.join(ROOT, '知识库.清单.json')
EXTS = ['.txt', '.md', '.markdown', '.csv', '.json', '.docx', '.pdf', '.html', '.htm']

PDF_SKIP = False       # --skip-pdf：完全不读 PDF（快）
PDF_PAGES = 0          # --pdf-pages N：每个 PDF 只读前 N 页
MAX_CHARS = 60000      # --max-chars N：每篇最多留 N 字
SPLIT_BOOKMARKS = True # --no-bookmarks：关闭 PDF 书签切分
OVERWRITE = False      # --overwrite：不合并、不读缓存，整份覆盖
DRY = False            # --dry：只统计不写


def read_text(path):
    for enc in ('utf-8-sig', 'utf-8', 'gbk', 'gb18030', 'big5'):
        try:
            with io.open(path, 'r', encoding=enc) as f:
                return f.read()
        except Exception:
            continue
    with io.open(path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()


def strip_html(s):
    s = re.sub(r'(?is)<(script|style)[^>]*>.*?</\1>', ' ', s)
    s = re.sub(r'(?s)<[^>]+>', ' ', s)
    s = s.replace('&nbsp;', ' ').replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&')
    return s


def flatten_json(o, out):
    if isinstance(o, dict):
        for v in o.values():
            if isinstance(v, (dict, list)):
                flatten_json(v, out)
            else:
                out.append(str(v))
    elif isinstance(o, list):
        for v in o:
            flatten_json(v, out)
    else:
        out.append(str(o))


def read_docx(path):
    try:
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(' | '.join(c.text.strip() for c in row.cells))
        txt = '\n'.join(parts)
        if txt.strip():
            return txt
    except Exception:
        pass
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read('word/document.xml').decode('utf-8', 'ignore')
        xml = xml.replace('</w:p>', '\n').replace('</w:tr>', '\n')
        return strip_html(xml)
    except Exception as e:
        print('  ! 读不了 docx：' + path + '（' + str(e) + '）')
        return ''


def _pdf_reader():
    try:
        from pypdf import PdfReader
        return PdfReader
    except Exception:
        try:
            from PyPDF2 import PdfReader
            return PdfReader
        except Exception:
            return None


def _pdf_outline_top(reader):
    """返回顶层书签 [(标题, 0 起始页码), ...]，按页码排序；嵌套子书签忽略（避免重复切分）。"""
    out = []
    try:
        items = reader.outline or []
    except Exception:
        return out
    for it in items:
        try:
            pg = reader.get_destination_page_number(it)
        except Exception:
            continue
        title = str(getattr(it, 'title', '') or '').strip()
        if title and isinstance(pg, int) and pg >= 0:
            out.append((title, pg))
    out.sort(key=lambda x: x[1])
    # 同一页有多个书签时只留第一个（避免切出空段）
    dedup, last = [], -1
    for t, p in out:
        if p != last:
            dedup.append((t, p))
            last = p
    return dedup


def page_from_filename(path, pages_total=None):
    """没书签时的兜底：有的精编版 PDF 是「按考点从原书里切出来的」，本身没有书签，
       但文件名里写着这一节在原书第几页 —— 如「1.结核_110-117.pdf」「1.感染性心内膜炎_34.pdf」。
       取到起始页就记下来（显示时就是用户翻书要找的那一页）。
       带范围时和 PDF 实际页数对一下：对不上就不采信（怕误伤别的文件名）。
       返回 1 基起始页，或 None。"""
    base = os.path.splitext(os.path.basename(path))[0]
    m = re.search(r'[_\-―－–—]\s*(\d{1,4})\s*(?:[-–—~～]\s*(\d{1,4}))?\s*$', base)
    if not m:
        return None
    a = int(m.group(1))
    b = int(m.group(2)) if m.group(2) else a
    if b < a:
        return None
    if pages_total and b > a and (b - a + 1) != pages_total:
        return None
    return a


def read_pdf(path):
    PdfReader = _pdf_reader()
    if PdfReader is None:
        print('  ! 读 PDF 需要 pypdf：在命令行里跑  pip install pypdf  再试')
        return []
    try:
        reader = PdfReader(path)
    except Exception as e:
        print('  ! 读不了 pdf：' + path + '（' + str(e) + '）')
        return []
    try:
        pages = reader.pages if not PDF_PAGES else reader.pages[:PDF_PAGES]
    except Exception:
        pages = list(reader.pages)
    page_texts = []
    for pg in pages:
        try:
            page_texts.append(pg.extract_text() or '')
        except Exception:
            page_texts.append('')
    if PDF_PAGES and len(reader.pages) > PDF_PAGES:
        print('    （只读了前 ' + str(PDF_PAGES) + ' 页 / 共 ' + str(len(reader.pages)) + ' 页）')
    pages_total = len(reader.pages)
    whole = '\n'.join(page_texts)
    if not whole.strip():
        return []

    if SPLIT_BOOKMARKS:
        bms = [(t, p) for t, p in _pdf_outline_top(reader) if p < len(page_texts)]
        if len(bms) >= 2:
            sections = []
            for i, (title, p) in enumerate(bms):
                end = bms[i + 1][1] if i + 1 < len(bms) else len(page_texts)
                txt = '\n'.join(page_texts[p:end]).strip()
                if txt:
                    sections.append({'name': title, 'text': txt, 'page': p + 1, 'pages': pages_total})
            # 第一个书签之前的封面 / 目录并进第一节，不丢内容（页码也归到 1）
            if bms[0][1] > 0:
                pre = '\n'.join(page_texts[:bms[0][1]]).strip()
                if pre and sections:
                    sections[0]['text'] = pre + '\n' + sections[0]['text']
                    sections[0]['page'] = 1
            if len(sections) >= 2:
                return sections
    fp = page_from_filename(path, pages_total)
    if fp:
        print('    （没有书签，按文件名记页码：第 ' + str(fp) + ' 页）')
    # pdfp：这一节在 PDF 文件里的第几页（文件名兜底时恒为 1 —— 文件本身就是从 fp 那页切出来的，
    # 直接拿 fp 当 #page= 会跳到文件最后一页）
    return [{'name': None, 'text': whole, 'page': fp, 'pages': None if fp else pages_total, 'pdfp': 1 if fp else None}]


def load_doc(path):
    """返回 [{name, text}, ...]；name 为 None 表示用相对路径做名字（普通文件）。PDF 书签切分返回多篇。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == '.docx':
        return [{'name': None, 'text': read_docx(path), 'page': None, 'pages': None}]
    if ext == '.pdf':
        return [] if PDF_SKIP else read_pdf(path)
    s = read_text(path) or ''
    if ext in ('.html', '.htm'):
        s = strip_html(s)
    if ext == '.json':
        try:
            out = []
            flatten_json(json.loads(s), out)
            s = '\n'.join(out)
        except Exception:
            pass
    return [{'name': None, 'text': s, 'page': None, 'pages': None}]


def clean(s):
    s = s.replace('\r\n', '\n').replace('\r', '\n')
    s = re.sub(r'[ \t]+\n', '\n', s)
    s = re.sub(r'\n{3,}', '\n\n', s)
    return s.strip()


def file_md5(path):
    h = hashlib.md5()
    try:
        with open(path, 'rb') as f:
            for chunk in iter(lambda: f.read(65536), b''):
                h.update(chunk)
    except Exception:
        return ''
    return h.hexdigest()


def scan(dirpath, old_by_file, old_md5):
    """扫一个目录。old_by_file / old_md5 用来做增量：没变的文件复用旧正文，不重新解析。
    返回 (docs, md5map)：docs 里每个 {name, text, file}，file = 相对路径（用来溯源）。"""
    docs = []
    md5map = {}
    for cur, dirs, files in os.walk(dirpath):
        dirs[:] = [d for d in dirs if not d.startswith('.')]
        for fn in sorted(files):
            ext = os.path.splitext(fn)[1].lower()
            if ext not in EXTS:
                continue
            if PDF_SKIP and ext == '.pdf':
                continue
            p = os.path.join(cur, fn)
            if os.path.abspath(p) == OUT:
                continue
            rel = os.path.relpath(p, dirpath)
            md5 = file_md5(p)
            md5map[rel] = md5
            # 增量复用：文件没变、且旧库里能拿到这段正文
            if not OVERWRITE and old_md5.get(rel) == md5 and rel in old_by_file:
                n = 0
                for d in old_by_file[rel]:
                    docs.append({'name': d['name'], 'text': d['text'], 'file': rel, 'page': d.get('page'), 'pages': d.get('pages'), 'pdfp': d.get('pdfp')})
                    n += 1
                if n:
                    print('  = ' + rel + '（没变，复用' + ('' if n == 1 else ' ' + str(n) + ' 节') + '）')
                    continue
            secs = load_doc(p)
            added = 0
            for s in secs:
                name = s['name'] or rel
                text = clean(s['text'])
                if MAX_CHARS and len(text) > MAX_CHARS:
                    text = text[:MAX_CHARS]
                if not text:
                    continue
                docs.append({'name': name, 'text': text, 'file': rel, 'page': s.get('page'), 'pages': s.get('pages'), 'pdfp': s.get('pdfp')})
                added += 1
            if added:
                chars = sum(len(d['text']) for d in docs[-added:])
                if added > 1:
                    print('  + ' + rel + '（书签切成 ' + str(added) + ' 节 / 共 ' + format(chars, ',') + ' 字）')
                else:
                    print('  + ' + rel + '（' + format(chars, ',') + ' 字）')
            else:
                print('  - 跳过（没读到文字）：' + fn)
    return docs, md5map


def parse_existing_kb(path):
    """把现有 知识库.js（一条或多条 window.__KBS__ = ...concat([...]);）解析回 libs 列表。"""
    s = read_text(path)
    if not s:
        return []
    libs = []
    dec = json.JSONDecoder()
    i = 0
    while True:
        i = s.find('concat(', i)
        if i < 0:
            break
        start = i + len('concat(')
        try:
            obj, end = dec.raw_decode(s, start)
        except Exception:
            i = start
            continue
        if isinstance(obj, list):
            libs.extend(obj)
        i = end
    return libs


def existing_by_file(lib):
    """把旧库的 docs 按「原文件路径」分组（file 字段优先，老格式退回 name），供增量复用。"""
    m = {}
    for d in lib.get('docs', []):
        f = d.get('file') or d.get('name') or ''
        m.setdefault(f, []).append({'name': d.get('name') or '', 'text': d.get('text') or '', 'page': d.get('page'), 'pages': d.get('pages'), 'pdfp': d.get('pdfp')})
    return m


def load_manifest():
    if OVERWRITE:
        return {}
    try:
        with io.open(MANIFEST, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


def save_manifest(manifest):
    if DRY:
        return
    with io.open(MANIFEST, 'w', encoding='utf-8') as f:
        json.dump(manifest, f, ensure_ascii=False)


def slug(name):
    return re.sub(r'[^\w\u4e00-\u9fa5-]+', '-', name).strip('-').lower()


USAGE = r"""刷题站 · 把文件夹转成知识库.js

用法：
    python make-kb.py                  扫站点目录里的 资料\ 文件夹
    python make-kb.py D:\书 D:\讲义     每个目录 = 一个知识库（可给多个）

可选参数：
    --overwrite      整份覆盖（不合并现有库、不读缓存，回到旧行为）
    --name 生理        给对应目录起个库名（按顺序，可重复：--name 外科 --name 内科）
    --skip-pdf       跳过 PDF（最快）
    --no-bookmarks   关闭 PDF 书签切分（一本书算一篇）
    --pdf-pages 30   每个 PDF 只读前 30 页
    --max-chars 60000 每篇最多留多少字
    --dry            只统计、不写文件
    -h, --help       看这段说明
"""


def main():
    global PDF_SKIP, PDF_PAGES, MAX_CHARS, SPLIT_BOOKMARKS, OVERWRITE, DRY
    raw = sys.argv[1:]
    args = []
    NAMES = []          # --name：给对应目录起库名（按顺序）
    i = 0
    while i < len(raw):
        a = raw[i]
        if a in ('-h', '--help'):
            print(USAGE)
            return
        if a == '--skip-pdf':
            PDF_SKIP = True
        elif a == '--name':
            i += 1
            if i < len(raw):
                NAMES.append(raw[i])
        elif a == '--no-bookmarks':
            SPLIT_BOOKMARKS = False
        elif a == '--overwrite':
            OVERWRITE = True
        elif a == '--dry':
            DRY = True
        elif a == '--pdf-pages':
            i += 1
            PDF_PAGES = int(raw[i]) if i < len(raw) and raw[i].isdigit() else 0
        elif a.startswith('--pdf-pages='):
            PDF_PAGES = int(a.split('=', 1)[1]) if a.split('=', 1)[1].isdigit() else 0
        elif a == '--max-chars':
            i += 1
            MAX_CHARS = int(raw[i]) if i < len(raw) and raw[i].isdigit() else MAX_CHARS
        elif a.startswith('--max-chars='):
            MAX_CHARS = int(a.split('=', 1)[1]) if a.split('=', 1)[1].isdigit() else MAX_CHARS
        elif a.startswith('-'):
            print('不认识的参数：' + a)
            print(USAGE)
            return
        else:
            args.append(a)
        i += 1
    if args:
        bases = []
        for idx, a in enumerate(args):
            ap = os.path.abspath(a)
            nm = NAMES[idx] if idx < len(NAMES) else (os.path.basename(ap) or a)
            bases.append((ap, nm))
    else:
        bases = [(os.path.join(ROOT, '资料'), '资料')]

    # 读现有库（合并模式）和增量清单
    existing_libs = {}
    if not OVERWRITE and os.path.exists(OUT):
        for lib in parse_existing_kb(OUT):
            if lib.get('id'):
                existing_libs[lib['id']] = lib
    manifest = load_manifest()

    final = dict(existing_libs)   # 先放现有库；被扫描的同名库会原位替换
    scanned = 0
    for base, name in bases:
        if not os.path.isdir(base):
            print('！目录不存在：' + base)
            continue
        print('扫描：' + base + ('　（合并模式，重跑只补差异）' if not OVERWRITE else '　（覆盖模式）'))
        lib_id = 'file-' + slug(name)
        old_lib = existing_libs.get(lib_id)
        old_by_file = existing_by_file(old_lib) if old_lib else {}
        old_md5 = manifest.get(lib_id, {})
        docs, md5map = scan(base, old_by_file, old_md5)
        if not docs:
            print('  （这个目录里没有可读的文档）')
            continue
        manifest[lib_id] = md5map
        final[lib_id] = {'id': lib_id, 'name': name, 'src': 'file', 'docs': docs}
        scanned += 1

    out_libs = list(final.values())
    if not out_libs:
        print('')
        print('没有生成内容。把 .txt / .md / .docx / .pdf 放进 资料\\ 文件夹，再双击一次 转知识库.bat。')
        return
    payload = json.dumps(out_libs, ensure_ascii=False).replace('</', '<\\/')
    total = sum(len(d['text']) for kb in out_libs for d in kb['docs'])
    ndoc = sum(len(kb['docs']) for kb in out_libs)
    print('')
    print('结果：' + str(len(out_libs)) + ' 个知识库 / ' + str(ndoc) + ' 篇 / 共 ' + format(total, ',') + ' 字'
          + ('　（本次扫描 ' + str(scanned) + ' 个目录）' if scanned else ''))
    if DRY:
        print('（--dry：只统计不写文件）')
        return
    with io.open(OUT, 'w', encoding='utf-8') as f:
        f.write('/* 由「转知识库.bat」自动生成。合并模式：重跑会按库 id 替换、保留其它库；别手改。 */\n')
        f.write('window.__KBS__ = (window.__KBS__ || []).concat(' + payload + ');\n')
    save_manifest(manifest)
    print('生成：' + OUT)
    if not OVERWRITE:
        print('增量清单：' + MANIFEST + '（可删，删了下次全量重解析）')
    print('回到刷题站按 Ctrl+F5 刷新，在「AI 与知识库」页面或 AI 面板的标签里就能看到它们。')


if __name__ == '__main__':
    main()
